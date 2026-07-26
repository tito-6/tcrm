# -*- coding: utf-8 -*-
# Part of TCRM. See LICENSE file for full copyright and licensing details.
from typing import Any, Dict, List, Optional

from .base import BaseMarketConnector, ConnectorError, FetchPageResult
from .registry import register


@register
class XlsxMarketConnector(BaseMarketConnector):
    source_type = "xlsx"
    display_name = "XLSX Upload"

    DEFAULT_CAPABILITIES = {
        **BaseMarketConnector.DEFAULT_CAPABILITIES,
        "pagination": False,
    }

    def validate_configuration(self):
        self.assert_enabled()
        return {"ok": True, "source_type": self.source_type, "needs_upload": True}

    def test_connection(self):
        self.assert_enabled()
        self.assert_authorized()
        return {"ok": True, "message": "XLSX connector ready for uploads."}

    def parse_bytes(self, data: bytes) -> List[Dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConnectorError("openpyxl is required for XLSX imports.") from exc
        import io
        wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip() for h in next(rows_iter, [])]
        if not headers:
            return []
        out = []
        for row in rows_iter:
            vals = list(row)
            if not any(v is not None and v != "" for v in vals):
                continue
            out.append({headers[i]: vals[i] if i < len(vals) else None for i in range(len(headers))})
        return out

    def fetch_page(self, cursor: Optional[str] = None, records=None, **kwargs) -> FetchPageResult:
        return FetchPageResult(records=records or [], next_cursor=None, exhausted=True)
